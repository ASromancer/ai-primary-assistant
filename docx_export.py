"""Xuất phiếu ra Word (.docx): phiếu chung, 3 phiếu nhóm, đáp án + ma trận; đóng gói ZIP."""
import io
import re
import zipfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ai import DANG_BAI, MUC, BanLuu, CauHoi

SAO = {1: "★", 2: "★★", 3: "★★★"}
# Phiếu nhóm: tên màu trung tính để học sinh không thấy "nhãn" mức độ.
# (mức chính, mức thử thách thêm)
NHOM = {
    "Xanh": ((1, 2), RGBColor(0x2E, 0x7D, 0x32), "HS cần hỗ trợ thêm"),
    "Cam": ((2, 3), RGBColor(0xEF, 0x6C, 0x00), "HS đại trà"),
    "Tím": ((3, None), RGBColor(0x6A, 0x1B, 0x9A), "HS hoàn thành tốt"),
}
XANH_DAM = RGBColor(0, 51, 102)
_TIEN_TO = re.compile(r"^\s*([0-9]{1,2}|[a-dA-D])\s*[.)]\s+")


def so(x: float) -> str:
    """Số kiểu Việt Nam: 1.5 -> '1,5'."""
    return f"{x:g}".replace(".", ",")


def bo_tien_to(s: str) -> str:
    """Bỏ 'A. ', '1) ' do AI tự thêm, để mình tự đánh số thống nhất."""
    return _TIEN_TO.sub("", s, count=1).strip()


def _doc_moi() -> Document:
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.5)
        s.left_margin, s.right_margin = Cm(2), Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(14)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(3)
    return doc


def _para(doc, text="", bold=False, italic=False, size=None, color=None, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if align is not None:
        p.alignment = align
    return p


def _tieu_de(doc, ts, truong, ten_phieu, mau=XANH_DAM):
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = f"Trường Tiểu học: {truong or '....................'}\nLớp: ............"
    t.cell(0, 1).text = "Họ và tên: ......................................\nNgày: ......./......./..........."
    _para(doc)
    _para(doc, ten_phieu, bold=True, size=16, color=mau, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, f"Môn {ts.mon} – Lớp {ts.lop} – Thời gian: {ts.thoi_luong} phút",
          italic=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, f"Chủ đề: {ts.chu_de}", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)


def _cau(doc, so_cau: int, c: CauHoi):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.add_run(f"Câu {so_cau} ({so(c.diem)} điểm). ").bold = True
    p.add_run(c.noi_dung)
    if c.goi_y_hs:
        _para(doc, f"Gợi ý: {c.goi_y_hs}", italic=True, size=12)
    ds = [bo_tien_to(x) for x in c.lua_chon]
    if c.dang == "trac_nghiem":
        for i, x in enumerate(ds):
            _para(doc, f"    {'ABCDEFGH'[i]}. {x}")
    elif c.dang == "dung_sai":
        t = doc.add_table(rows=1, cols=2, style="Table Grid")
        t.autofit = False
        t.cell(0, 0).text, t.cell(0, 1).text = "Nhận định", "Đ / S"
        for x in ds:
            row = t.add_row().cells
            row[0].text = x
        for row in t.rows:
            row.cells[0].width, row.cells[1].width = Cm(14), Cm(2.5)
    elif c.dang == "noi_cot":
        phai = [bo_tien_to(x) for x in c.cot_phai]
        t = doc.add_table(rows=max(len(ds), len(phai)), cols=3)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, x in enumerate(ds):
            t.cell(i, 0).text = f"{i + 1}. {x}   •"
        for i, x in enumerate(phai):
            t.cell(i, 2).text = f"•   {'abcdefgh'[i]}. {x}"
    elif c.dang == "tu_luan":
        for _ in range(c.muc + 2):
            _para(doc, "." * 140, size=12)


def _cuoi_phieu(doc, ban: BanLuu):
    _para(doc)
    _para(doc, ban.phieu.loi_chuc, italic=True, color=XANH_DAM, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, "Em tự đánh giá:   ☐ Em làm được hết     ☐ Em làm được một phần     ☐ Em cần cô giúp thêm", size=12)
    _para(doc, "Nhận xét của giáo viên: " + "." * 105, size=12)
    _para(doc, "." * 140, size=12)


def _luu(doc) -> bytes:
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def cau_theo_muc(ban: BanLuu, muc) -> list[int]:
    return [i for i, c in enumerate(ban.phieu.cau_hoi) if c.muc == muc]


def cau_cua_nhom(ban: BanLuu, ten: str) -> tuple[list[int], list[int]]:
    """(câu chính, câu thử thách thêm) của phiếu nhóm: mọi câu mức chính + câu đầu của mức kế tiếp."""
    (chinh, them), _, _ = NHOM[ten]
    return cau_theo_muc(ban, chinh), cau_theo_muc(ban, them)[:1]


def phieu_chung(ban: BanLuu) -> bytes:
    doc = _doc_moi()
    _tieu_de(doc, ban.thong_so, ban.thong_so.truong, "PHIẾU HỌC TẬP")
    stt = 0
    for muc in (1, 2, 3):
        ids = cau_theo_muc(ban, muc)
        if ids:
            _para(doc, f"{SAO[muc]}  THỬ THÁCH {muc} SAO", bold=True, color=XANH_DAM).paragraph_format.space_before = Pt(10)
        for i in ids:
            stt += 1
            _cau(doc, stt, ban.phieu.cau_hoi[i])
    _cuoi_phieu(doc, ban)
    return _luu(doc)


def phieu_nhom(ban: BanLuu, ten: str) -> bytes:
    _, mau, _ = NHOM[ten]
    chinh, them = cau_cua_nhom(ban, ten)
    doc = _doc_moi()
    _tieu_de(doc, ban.thong_so, ban.thong_so.truong, f"PHIẾU HỌC TẬP – PHIẾU {ten.upper()}", mau)
    for stt, i in enumerate(chinh, 1):
        _cau(doc, stt, ban.phieu.cau_hoi[i])
    if them:
        _para(doc, "★ THỬ THÁCH THÊM (nếu em còn thời gian)", bold=True, color=mau).paragraph_format.space_before = Pt(10)
        _cau(doc, len(chinh) + 1, ban.phieu.cau_hoi[them[0]])
    _cuoi_phieu(doc, ban)
    return _luu(doc)


def vi_tri_trong_phieu(ban: BanLuu) -> dict[int, str]:
    """index câu -> 'Chung: 3; Cam: 1' để GV tra đáp án cho mọi phiếu."""
    thu_tu = [i for m in (1, 2, 3) for i in cau_theo_muc(ban, m)]
    out = {i: [f"Chung: {k + 1}"] for k, i in enumerate(thu_tu)}
    for ten in NHOM:
        chinh, them = cau_cua_nhom(ban, ten)
        for k, i in enumerate(chinh + them):
            out[i].append(f"{ten}: {k + 1}")
    return {i: "; ".join(v) for i, v in out.items()}


def ma_tran(ban: BanLuu) -> list[list]:
    """Hàng: mức 1..3 + Tổng. Cột: Mức | từng dạng bài | Số câu | Điểm."""
    rows = []
    for muc in (1, 2, 3, None):
        cs = [c for c in ban.phieu.cau_hoi if muc is None or c.muc == muc]
        rows.append([f"Mức {muc}" if muc else "Tổng"]
                    + [sum(c.dang == d for c in cs) for d in DANG_BAI]
                    + [len(cs), round(sum(c.diem for c in cs), 2)])
    return rows


def dap_an(ban: BanLuu) -> bytes:
    ts, phieu = ban.thong_so, ban.phieu
    doc = _doc_moi()
    _para(doc, "ĐÁP ÁN – MA TRẬN ĐỀ (dành cho giáo viên)", bold=True, size=16, color=XANH_DAM,
          align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, f"Môn {ts.mon} – Lớp {ts.lop} – Bộ sách {ts.bo_sach} – {ts.thoi_luong} phút\nChủ đề: {ts.chu_de}",
          italic=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)

    _para(doc, "1. Yêu cầu cần đạt", bold=True)
    for y in phieu.yeu_cau_can_dat:
        _para(doc, f"– {y}", size=13)

    _para(doc, "2. Ma trận theo mức độ (Thông tư 27/2020/TT-BGDĐT)", bold=True)
    head = ["Mức"] + list(DANG_BAI.values()) + ["Số câu", "Điểm"]
    data = ma_tran(ban)
    t = doc.add_table(rows=1 + len(data), cols=len(head), style="Table Grid")
    for j, h in enumerate(head):
        t.cell(0, j).text = h
    for i, row in enumerate(data, 1):
        for j, v in enumerate(row):
            t.cell(i, j).text = so(v) if isinstance(v, float) else str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(11)
    for muc, mo_ta in MUC.items():
        _para(doc, f"Mức {muc}: {mo_ta}.", italic=True, size=11)

    _para(doc, "3. Phân phiếu theo nhóm", bold=True)
    for ten, ((chinh, them), _, nhom_hs) in NHOM.items():
        _para(doc, f"– Phiếu {ten}: {nhom_hs} (câu Mức {chinh}" + (f" + 1 câu thử thách Mức {them})" if them else ")"),
              size=13)

    _para(doc, "4. Đáp án và hướng dẫn chấm", bold=True)
    vt = vi_tri_trong_phieu(ban)
    t = doc.add_table(rows=1, cols=4, style="Table Grid")
    for j, h in enumerate(["Câu (vị trí trong phiếu)", "Mức – Dạng", "Đáp án", "Hướng dẫn chấm"]):
        t.cell(0, j).text = h
    for i in (i for m in (1, 2, 3) for i in cau_theo_muc(ban, m)):
        c = phieu.cau_hoi[i]
        r = t.add_row().cells
        r[0].text = vt[i]
        r[1].text = f"Mức {c.muc} – {DANG_BAI[c.dang]} ({so(c.diem)}đ)"
        r[2].text, r[3].text = c.dap_an, c.huong_dan_cham
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(12)
    return _luu(doc)


def tat_ca(ban: BanLuu) -> dict[str, bytes]:
    files = {"1_Phieu_chung.docx": phieu_chung(ban)}
    for k, ten in enumerate(NHOM, 2):
        files[f"{k}_Phieu_{ten.replace('í', 'i')}.docx"] = phieu_nhom(ban, ten)
    files["5_Dap_an_Ma_tran.docx"] = dap_an(ban)
    return files


def dong_goi_zip(files: dict[str, bytes]) -> bytes:
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z:
        for ten, data in files.items():
            z.writestr(ten, data)
    return bio.getvalue()
