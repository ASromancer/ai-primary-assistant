"""Đánh giá theo TT27 từ kết quả chấm: mức đạt, khớp tên học sinh, sổ tổng hợp Excel."""
import difflib
import io
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

TY_LE = {"dung": 1.0, "mot_phan": 0.5, "sai": 0.0, "bo_trong": 0.0}
KY_HIEU = {"dung": "✓", "mot_phan": "½", "sai": "✗", "bo_trong": "–"}
DAO_KY_HIEU = {v: k for k, v in KY_HIEU.items()}
MUC_TT27 = ["Hoàn thành tốt", "Hoàn thành", "Chưa hoàn thành"]
NGUONG_BO_TRO = 0.4  # câu có từ 40% học sinh chưa đạt trở lên cần dạy bổ trợ
NHOM_DE_XUAT = {"Hoàn thành tốt": "Tím", "Hoàn thành": "Cam", "Chưa hoàn thành": "Xanh"}


def ty_le_theo_muc(phieu, ids: list[int], kq) -> dict[int, float | None]:
    """Tỉ lệ làm đúng từng mức trên phiếu đã phát (ids: chỉ số câu theo thứ tự trên phiếu)."""
    dat = {c.cau_so: TY_LE[c.dat] for c in kq.cau}
    out = {}
    for muc in (1, 2, 3):
        diem = [dat.get(k, 0.0) for k, i in enumerate(ids, 1) if phieu.cau_hoi[i].muc == muc]
        out[muc] = sum(diem) / len(diem) if diem else None
    return out


def muc_tt27(phieu, ids: list[int], kq) -> str:
    """Quy tắc cố định, giải thích được (mức không có câu coi như đạt):
    CHT: Mức 1 < 50%. HTT: Mức 1 ≥ 80%, Mức 2 ≥ 80%, Mức 3 ≥ 50%. Còn lại: HT."""
    t = ty_le_theo_muc(phieu, ids, kq)
    dat = lambda muc, nguong: t[muc] is None or t[muc] >= nguong
    if not dat(1, 0.5):
        return "Chưa hoàn thành"
    if dat(1, 0.8) and dat(2, 0.8) and dat(3, 0.5):
        return "Hoàn thành tốt"
    return "Hoàn thành"


def muc_cuoi_ky(ds_muc: list[str]) -> str:
    """Mức gợi ý cuối kỳ: mức xuất hiện nhiều nhất; hoà thì lấy mức của lần gần nhất trong số các mức hoà."""
    if not ds_muc:
        return ""
    dem = {m: ds_muc.count(m) for m in ds_muc}
    cao = max(dem.values())
    return next(m for m in reversed(ds_muc) if dem[m] == cao)


def phan_tich_cau(rows: list[dict], so_cau: int) -> list[dict]:
    """Thống kê từng câu từ sổ tổng hợp (ô dạng '✓', '½ ghi chú', '✗ ghi chú', '–').
    chua_dat = (sai + bỏ trống + ½ một phần) / số bài; loi = các ghi chú lỗi khác nhau."""
    out = []
    for k in range(1, so_cau + 1):
        dem = dict.fromkeys(KY_HIEU, 0)
        loi = []
        for r in rows:
            o = str(r.get(f"Câu {k}") or "").strip()
            dat = DAO_KY_HIEU.get(o[:1])
            if dat:
                dem[dat] += 1
            if o[1:].strip() and o[1:].strip() not in loi:
                loi.append(o[1:].strip())
        n = sum(dem.values()) or 1
        out.append({"cau": k, **dem, "chua_dat": (dem["sai"] + dem["bo_trong"] + 0.5 * dem["mot_phan"]) / n,
                    "loi": loi[:8]})
    return out


def yeu_cau_bo_tro(cau_hoi: list, phan_tich: list[dict], toi_da: int = 4) -> tuple[str, list[int]]:
    """Chọn câu yếu (≥ ngưỡng; nếu không có thì câu yếu nhất còn lỗi) và dựng yêu cầu cho AI soạn phiếu bổ trợ.
    Chỉ dùng nội dung câu hỏi và ghi chú lỗi – không có tên học sinh."""
    xep = sorted(phan_tich, key=lambda p: -p["chua_dat"])
    chon = [p for p in xep if p["chua_dat"] >= NGUONG_BO_TRO][:toi_da] or [p for p in xep if p["chua_dat"] > 0][:1]
    dong = []
    for p in chon:
        c = cau_hoi[p["cau"] - 1]
        dong.append(f"- Câu {p['cau']} (Mức {c.muc}): \"{c.noi_dung}\" – đáp án: {c.dap_an}. "
                    f"{p['chua_dat']:.0%} học sinh chưa đạt" + (f"; lỗi thường gặp: {'; '.join(p['loi'])}" if p["loi"] else ""))
    yc = ("Đây là PHIẾU BỔ TRỢ soạn sau khi chấm bài. Học sinh còn yếu ở các nội dung sau:\n" + "\n".join(dong) +
          "\nHãy tập trung luyện lại đúng các kiến thức, kỹ năng này (không lặp nguyên câu cũ): Mức 1 chia nhỏ từng "
          "bước, giúp tránh lỗi thường gặp; Mức 2, 3 vận dụng lại trong tình huống mới.")
    return yc, [p["cau"] for p in chon]


def bo_dau(s: str) -> str:
    s = s.replace("đ", "d").replace("Đ", "D")
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()


def khop_ten(ten: str, ds_ten: list[str]) -> str:
    """Khớp tên AI đọc được trên phiếu với danh sách lớp (bỏ dấu, không phân biệt hoa thường)."""
    chuan = lambda s: " ".join(bo_dau(s).lower().split())
    if not ten.strip():
        return ""
    bang = {chuan(x): x for x in ds_ten}
    gan = difflib.get_close_matches(chuan(ten), list(bang), n=1, cutoff=0.6)
    return bang[gan[0]] if gan else ten.strip()


def so_tong_hop_excel(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tổng hợp"
    cot = list(rows[0]) if rows else []
    ws.append(cot)
    for r in rows:
        ws.append([r.get(c) for c in cot])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1565C0")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for k, c in enumerate(cot, 1):
        rong = max([len(str(c))] + [len(str(r.get(c) or "")) for r in rows])
        ws.column_dimensions[ws.cell(1, k).column_letter].width = min(60, max(6, rong + 2))
        if c == "Nhận xét":
            for cell in ws.iter_rows(min_row=2, min_col=k, max_col=k):
                cell[0].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ---------------- Tự chấm bài làm online ----------------
_CHO_TRONG = re.compile(r"\.{3,}|…+|_{3,}")
_DUNG_SAI = re.compile(r"(?<![A-Za-zÀ-ỹĐđ])(Đúng|Sai|Đ|S)(?![A-Za-zÀ-ỹĐđ])")
_NOI = re.compile(r"(\d+)\s*[-–—:→>]+\s*([a-hA-H])\b")
CHU = "ABCDEFGH"


def so_cho_trong(noi_dung: str) -> int:
    """Số chỗ trống ('.......', '…', '___') trong đề điền khuyết; tối thiểu 1."""
    return max(1, len(_CHO_TRONG.findall(noi_dung)))


def dap_an_dung_sai(c) -> list[bool] | None:
    """Đọc đáp án dạng 'a) Đ; b) S; c) Đ'. None nếu không khớp số nhận định (để AI chấm)."""
    kq = [m in ("Đ", "Đúng") for m in _DUNG_SAI.findall(c.dap_an)]
    return kq if c.lua_chon and len(kq) == len(c.lua_chon) else None


def dap_an_noi_cot(c) -> dict[int, int] | None:
    """Đọc đáp án dạng '1 – b; 2 – c' thành {chỉ số trái: chỉ số phải}. None nếu không đủ cặp."""
    cap = {int(i) - 1: CHU.lower().index(ch.lower()) for i, ch in _NOI.findall(c.dap_an)}
    return cap if c.lua_chon and len(cap) == len(c.lua_chon) else None


def _ket_qua(c, dung: int, tong: int, bo_trong: bool):
    from ai import KetQuaCau
    if bo_trong:
        return KetQuaCau(cau_so=0, dat="bo_trong", diem_dat=0, ghi_chu="")
    ty_le = dung / tong
    dat = "dung" if ty_le == 1 else "sai" if ty_le == 0 else "mot_phan"
    return KetQuaCau(cau_so=0, dat=dat, diem_dat=round(c.diem * ty_le, 2), ghi_chu="")


def tu_cham(c, tra_loi):
    """Chấm bằng code các câu trắc nghiệm / đúng-sai / nối cột (chấm điểm từng ý).
    Trả None nếu cần AI (điền khuyết, tự luận, hoặc không đọc được đáp án)."""
    if c.dang == "trac_nghiem":
        from classroom_html import dap_an_dung
        k = dap_an_dung(c)
        return None if k is None else _ket_qua(c, int(tra_loi == k), 1, tra_loi is None)
    if c.dang == "dung_sai":
        dung = dap_an_dung_sai(c)
        tl = list(tra_loi or [None] * len(c.lua_chon))
        return None if dung is None else _ket_qua(c, sum(a == b for a, b in zip(tl, dung)), len(dung),
                                                  all(x is None for x in tl))
    if c.dang == "noi_cot":
        cap = dap_an_noi_cot(c)
        tl = list(tra_loi or [None] * len(c.lua_chon))
        return None if cap is None else _ket_qua(c, sum(tl[i] == j for i, j in cap.items() if i < len(tl)), len(cap),
                                                 all(x is None for x in tl))
    return None


def hien_thi_tra_loi(c, tra_loi) -> str:
    """Câu trả lời của học sinh dạng chữ (cho AI chấm và bảng của giáo viên)."""
    if c.dang == "trac_nghiem":
        from docx_export import bo_tien_to
        return f"{CHU[tra_loi]}. {bo_tien_to(c.lua_chon[tra_loi])}" if tra_loi is not None and tra_loi < len(c.lua_chon) else "–"
    if c.dang == "dung_sai":
        return "; ".join(f"{CHU[i].lower()}) {'Đ' if v else '–' if v is None else 'S'}"
                         for i, v in enumerate(tra_loi or []))
    if c.dang == "noi_cot":
        return "; ".join(f"{i + 1}–{CHU[v].lower() if v is not None else '?'}" for i, v in enumerate(tra_loi or []))
    if c.dang == "dien_khuyet":
        return " | ".join((x or "").strip() or "…" for x in (tra_loi or []))
    return (tra_loi or "").strip() or "–"
