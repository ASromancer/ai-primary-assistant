"""Chạy: python test_app.py (không cần API key, không gọi mạng)."""
import io
import zipfile
from types import SimpleNamespace

from docx import Document
from streamlit.testing.v1 import AppTest

import ai
import classroom_html as cx
import danh_gia as dg
import docx_export as dx

BAN = ai.BanLuu.model_validate_json(open("mau/phieu_mau.json", encoding="utf-8").read())


class FakeClient:
    """Trả lần lượt các kết quả parsed định sẵn, thay cho Gemini."""
    def __init__(self, *parsed):
        self.parsed, self.calls = list(parsed), 0
        self.models = self

    def generate_content(self, model, contents, config):
        self.calls += 1
        return SimpleNamespace(parsed=self.parsed.pop(0))


def test_tao_phieu_sap_xep_va_thu_lai():
    p = BAN.phieu.model_copy(deep=True)
    p.cau_hoi.reverse()
    p.cau_hoi[0].muc = 7  # AI trả mức sai -> kẹp về 3
    client = FakeClient(None, p)  # lần 1 JSON hỏng -> thử lại
    out = ai.tao_phieu(client, "m", BAN.thong_so, list(ai.DANG_BAI), [])
    assert client.calls == 2
    assert [c.muc for c in out.cau_hoi] == sorted(c.muc for c in out.cau_hoi)
    assert max(c.muc for c in out.cau_hoi) == 3


def test_chuyen_model_du_phong_khi_qua_tai():
    class Client503(FakeClient):
        def generate_content(self, model, contents, config):
            self.models_used.append(model)
            if model == "m":
                raise ai.errors.APIError(503, {"error": {"message": "overloaded", "status": "UNAVAILABLE"}})
            return super().generate_content(model, contents, config)
    client = Client503(BAN.phieu.model_copy(deep=True))
    client.models_used = []
    ai.tao_phieu(client, "m", BAN.thong_so, list(ai.DANG_BAI), [])
    assert client.models_used == ["m", ai.MODEL_DU_PHONG[0]]

    sai_key = FakeClient()
    sai_key.generate_content = lambda **kw: (_ for _ in ()).throw(ai.errors.APIError(400, {"error": {"message": "API key not valid"}}))
    try:
        ai.tao_phieu(sai_key, "m", BAN.thong_so, list(ai.DANG_BAI), [])
        assert False, "lỗi 400 phải báo ngay, không chuyển model"
    except ai.errors.APIError as e:
        assert e.code == 400


def test_tao_lai_cau_giu_muc_va_diem():
    moi = BAN.phieu.cau_hoi[0].model_copy(update=dict(muc=3, diem=5, noi_dung="Câu mới"))
    out = ai.tao_lai_cau(FakeClient(moi), "m", BAN.thong_so, BAN.phieu, 0, "de_hon")
    assert (out.muc, out.diem, out.noi_dung) == (1, 1, "Câu mới")


def test_kiem_dinh_kep_diem_va_loc_cau():
    kd = ai.KiemDinh(diem=140, nhan_xet_chung="ok", van_de=[
        ai.VanDe(cau_so=2, loai="dap_an_sai", mo_ta="sai", de_xuat="sửa"),
        ai.VanDe(cau_so=99, loai="ngon_ngu", mo_ta="x", de_xuat="y")])
    out = ai.kiem_dinh(FakeClient(kd), "m", BAN.thong_so, BAN.phieu)
    assert out.diem == 100 and [v.cau_so for v in out.van_de] == [2]


def test_dau_van_tay_doi_khi_sua():
    p = BAN.phieu.model_copy(deep=True)
    a = ai.dau_van_tay(p)
    p.cau_hoi[0].noi_dung += "!"
    assert a != ai.dau_van_tay(p)


def test_xuat_word():
    files = dx.tat_ca(BAN)
    assert len(files) == 5
    with zipfile.ZipFile(io.BytesIO(dx.dong_goi_zip(files))) as z:
        assert sorted(z.namelist()) == sorted(files)

    def text(b):
        d = Document(io.BytesIO(b))
        return "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])

    chung = text(files["1_Phieu_chung.docx"])
    assert all(c.noi_dung in chung for c in BAN.phieu.cau_hoi)
    assert "Mức" not in chung  # phiếu học sinh không lộ nhãn mức độ
    xanh = text(files["2_Phieu_Xanh.docx"])  # 3 câu mức 1 + 1 câu thử thách mức 2
    assert "Câu 4" in xanh and "Câu 5" not in xanh and "THỬ THÁCH THÊM" in xanh
    assert "THỬ THÁCH THÊM" not in text(files["4_Phieu_Tim.docx"])
    da = text(files["5_Dap_an_Ma_tran.docx"])
    assert "B. 43" in da and "Chung: 1; Xanh: 1" in da
    assert dx.ma_tran(BAN)[-1][-2:] == [7, 10.0]


def test_trang_in_html():
    import html_export as hx
    ban = BAN.model_copy(deep=True)
    ban.phieu.cau_hoi[0].noi_dung = "<script>alert(1)</script> 38 + 5"
    trang = hx.cac_trang(ban)
    assert list(trang) == ["Phiếu chung", "Phiếu Xanh", "Phiếu Cam", "Phiếu Tím", "Đáp án & ma trận"]
    doc = hx.trang_in(list(trang.values()), 18)
    assert "<script>alert" not in doc and "&lt;script&gt;" in doc  # nội dung AI phải được escape
    assert "font-size: 18pt" in doc and doc.count('class="trang') == 5
    assert "Mức" not in trang["Phiếu chung"] and "THỬ THÁCH THÊM" in trang["Phiếu Xanh"]


def test_dap_an_dung():
    c = BAN.phieu.cau_hoi[0]
    assert cx.dap_an_dung(c) == 1  # "B. 43"
    assert cx.dap_an_dung(c.model_copy(update={"dap_an": "Đáp án: C"})) == 2
    assert cx.dap_an_dung(c.model_copy(update={"dap_an": "53"})) == 3  # khớp nội dung phương án
    assert cx.dap_an_dung(c.model_copy(update={"dap_an": "Không rõ"})) is None
    assert cx.dap_an_dung(BAN.phieu.cau_hoi[1]) is None  # không phải trắc nghiệm


def test_trinh_chieu_an_toan():
    ban = BAN.model_copy(deep=True)
    ban.phieu.cau_hoi[0].noi_dung = "</script><script>alert(1)</script>"
    h = cx.trinh_chieu(ban, ["An", "Bình"], 60, True)
    assert "</script><script>alert" not in h and "<\\/script>" in h
    assert '"giay": 60' in h and "speechSynthesis" in h and '"An"' in h


def test_de_b_giu_cau_truc():
    b = BAN.phieu.model_copy(deep=True)
    for c in b.cau_hoi:
        c.muc, c.diem = 3, 9
    out = ai.tao_de_b(FakeClient(b), "m", BAN.thong_so, BAN.phieu)
    assert [(c.muc, c.dang, c.diem) for c in out.cau_hoi] == [(c.muc, c.dang, c.diem) for c in BAN.phieu.cau_hoi]
    thieu = BAN.phieu.model_copy(update={"cau_hoi": BAN.phieu.cau_hoi[:2]})
    try:
        ai.tao_de_b(FakeClient(thieu, thieu), "m", BAN.thong_so, BAN.phieu)
        assert False, "số câu khác phải báo lỗi"
    except ValueError:
        pass


def test_phieu_theo_ten():
    import html_export as hx
    ban = BAN.model_copy(update={"phieu_b": BAN.phieu.model_copy(deep=True)})
    hs = [ai.HocSinh(stt=1, ten="Nguyễn An", nhom="Xanh"), ai.HocSinh(stt=2, ten="Lê Bình", nhom="Tím")]
    trang = hx.phieu_theo_ten(ban, hs, True)
    assert len(trang) == 2 and "Nguyễn An" in trang[0] and "PHIẾU XANH" in trang[0]
    assert "ĐỀ B" in trang[1] and "ĐỀ B" not in trang[0]
    assert "ĐỀ B" not in "".join(hx.phieu_theo_ten(ban, hs, False))
    d = Document(io.BytesIO(dx.phieu_theo_ten(ban, hs, True)))
    text = "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])
    assert "Nguyễn An" in text and "Lê Bình" in text and "PHIẾU TÍM – ĐỀ B" in text


def _kq(*dat):
    return ai.KetQuaCham(ten_tren_phieu="An", nhan_xet="", cau=[
        ai.KetQuaCau(cau_so=k + 1, dat=d, diem_dat=0, ghi_chu="") for k, d in enumerate(dat)])


def test_muc_tt27():
    ids = list(range(7))  # 3 câu mức 1, 2 câu mức 2, 2 câu mức 3
    assert dg.muc_tt27(BAN.phieu, ids, _kq(*["dung"] * 7)) == "Hoàn thành tốt"
    assert dg.muc_tt27(BAN.phieu, ids, _kq("sai", "sai", "dung", "dung", "dung", "dung", "dung")) == "Chưa hoàn thành"
    assert dg.muc_tt27(BAN.phieu, ids, _kq("dung", "dung", "dung", "dung", "mot_phan", "sai", "sai")) == "Hoàn thành"
    assert dg.muc_tt27(BAN.phieu, ids, _kq("dung", "dung", "dung", "dung", "dung", "mot_phan", "sai")) == "Hoàn thành"
    xanh = [0, 1, 2, 3]  # phiếu Xanh: không có mức 3 -> mức 3 coi như đạt
    assert dg.muc_tt27(BAN.phieu, xanh, _kq("dung", "dung", "dung", "dung")) == "Hoàn thành tốt"


def test_khop_ten():
    assert dg.khop_ten("nguyen van an", ["Nguyễn Văn An", "Lê Bình"]) == "Nguyễn Văn An"
    assert dg.khop_ten("Le Binh", ["Nguyễn Văn An", "Lê Bình"]) == "Lê Bình"
    assert dg.khop_ten("Trần C", ["Nguyễn Văn An"]) == "Trần C"
    assert dg.khop_ten("", ["Nguyễn Văn An"]) == ""


def test_excel():
    from openpyxl import load_workbook
    b = dg.so_tong_hop_excel([{"STT": 1, "Họ tên": "An", "Điểm": 9.5, "Mức TT27": "Hoàn thành tốt", "Nhận xét": "Tốt"}])
    ws = load_workbook(io.BytesIO(b)).active
    assert ws["B2"].value == "An" and ws["A1"].font.bold and ws["C2"].value == 9.5


def test_cham_bai_kep_diem_va_bu_cau():
    kq = _kq("dung")
    kq.cau[0].diem_dat = 99
    out = ai.cham_bai(FakeClient(kq), "m", BAN.thong_so, BAN.phieu, [0, 1], (b"x", "image/png"))
    assert out.cau[0].diem_dat == BAN.phieu.cau_hoi[0].diem
    assert [c.cau_so for c in out.cau] == [1, 2] and out.cau[1].dat == "bo_trong"


def _sqlite():
    import tempfile
    from sqlalchemy import create_engine
    import kho
    db = create_engine(f"sqlite:///{tempfile.mkdtemp()}/t.db")
    kho.tao_bang(db)
    return db


def kiem_tra_kho(db):
    import kho
    A, B = "a@truong.vn", "b@truong.vn"
    kho.ho_so_luu(db, A, ten="Cô Lan", lop_mac_dinh=2, khong_hop_le="x")
    assert kho.ho_so_lay(db, A)["ten"] == "Cô Lan" and kho.ho_so_lay(db, B) == {}

    ban = BAN.model_copy(deep=True)
    pid = kho.phieu_luu(db, A, ban)
    assert ban.id == pid and kho.phieu_luu(db, A, ban) == pid  # cập nhật, không tạo mới
    ds = kho.phieu_ds(db, A)
    assert len(ds) == 1 and ds[0]["diem_kd"] == 92 and ds[0]["so_cau"] == 7
    kho.phieu_sao(db, A, pid, True)
    assert kho.phieu_ds(db, A)[0]["gan_sao"] is True
    # Cách ly: B không mở/sửa/xoá/gắn sao được phiếu của A; lưu với id của A tạo bản riêng cho B
    assert kho.phieu_mo(db, B, pid) is None and kho.phieu_ds(db, B) == []
    kho.phieu_xoa(db, B, pid); kho.phieu_sao(db, B, pid, False)
    assert kho.phieu_ds(db, A)[0]["gan_sao"] is True
    gia = BAN.model_copy(update={"id": pid})
    assert kho.phieu_luu(db, B, gia) != pid and kho.phieu_mo(db, A, pid).thong_so.chu_de == BAN.thong_so.chu_de

    lop = ai.LopHoc(ten_lop="2A", hoc_sinh=[ai.HocSinh(stt=1, ten="An", nhom="Xanh"), ai.HocSinh(stt=2, ten="Bình")])
    lid = kho.lop_luu(db, A, lop)
    lop.hoc_sinh = lop.hoc_sinh[:1]
    kho.lop_luu(db, A, lop)
    assert kho.lop_mo(db, A, lid).hoc_sinh == [ai.HocSinh(stt=1, ten="An", nhom="Xanh")]
    assert kho.lop_ds(db, A)[0]["si_so"] == 1
    lop_gia = ai.LopHoc(id=lid, ten_lop="hack", hoc_sinh=[])
    kho.lop_luu(db, B, lop_gia)  # không được xoá học sinh lớp của A
    assert lop_gia.id != lid and kho.lop_mo(db, A, lid).ten_lop == "2A" and kho.lop_mo(db, B, lid) is None

    rows = [{"ten": "An", "diem": 9, "muc_tt27": "Hoàn thành tốt", "nhan_xet": "Tốt", "chi_tiet": {"Câu 1": "✓"}}]
    kho.cham_luu(db, A, lid, "Toán", "Phép cộng", "Phiếu chung", rows)
    kho.cham_luu(db, A, lid, "Toán", "Phép trừ", "Phiếu chung", [{**rows[0], "muc_tt27": "Hoàn thành"}])
    kho.cham_luu(db, B, lid, "Toán", "x", "x", rows)  # B không gắn được vào lớp của A
    tb = kho.tien_bo(db, A, lid)
    assert [r["muc_tt27"] for r in tb] == ["Hoàn thành tốt", "Hoàn thành"] and kho.tien_bo(db, B, lid) == []
    tk = kho.thong_ke(db, A)
    assert (tk["so_phieu"], tk["so_lop"], tk["so_hoc_sinh"], tk["so_bai_cham"]) == (1, 1, 1, 2)
    kho.lop_xoa(db, B, lid)
    assert kho.lop_mo(db, A, lid) is not None
    kho.lop_xoa(db, A, lid)
    assert kho.lop_ds(db, A) == [] and kho.thong_ke(db, A)["so_bai_cham"] == 2  # lịch sử chấm vẫn giữ


def test_kho_sqlite():
    kiem_tra_kho(_sqlite())


def test_kho_postgres():
    import os
    url = os.environ.get("TEST_PG_URL")
    if not url:
        print("  (bỏ qua: chưa đặt TEST_PG_URL)")
        return
    from sqlalchemy import create_engine
    import kho
    db = create_engine(url)
    kho.tao_bang(db, postgres=True)
    kiem_tra_kho(db)


def test_bo_tien_to():
    assert dx.bo_tien_to("A. 43") == "43"
    assert dx.bo_tien_to("2) quả cam") == "quả cam"
    assert dx.bo_tien_to("1.5 kg") == "1.5 kg"


def _app(secret_key=None):
    at = AppTest.from_file("app.py", default_timeout=30)
    if secret_key:
        at.secrets["GEMINI_API_KEY"] = secret_key
    return at.run()


def test_trang_chu_va_dieu_huong():
    at = _app()
    assert not at.exception
    for trang in ["trang/soan_phieu.py", "trang/lop_hoc.py", "trang/cham_bai.py", "trang/lop_cua_toi.py"]:
        at.switch_page(trang).run()
        assert not at.exception, trang


def test_soan_phieu_mau_doi_cau_hoan_tac():
    goc = ai.tao_lai_cau
    ai.tao_lai_cau = lambda client, model, ts, phieu, i, che_do, **kw: phieu.cau_hoi[i].model_copy(update={"noi_dung": "CÂU MỚI"})
    try:
        at = _app("fake")
        at.switch_page("trang/soan_phieu.py").run()
        at.button(key="mau_main").click().run()
        assert not at.exception
        assert len(at.session_state.lich_su) == 1
        cu = at.session_state.ban.phieu.cau_hoi[0].noi_dung
        at.button(key="doi0").click().run()
        assert not at.exception and at.session_state.ban.phieu.cau_hoi[0].noi_dung == "CÂU MỚI"
        at.button(key="ht0").click().run()
        assert at.session_state.ban.phieu.cau_hoi[0].noi_dung == cu
    finally:
        ai.tao_lai_cau = goc


def test_kiem_dinh_tren_giao_dien():
    goc = ai.kiem_dinh, ai.tao_lai_cau
    ai.kiem_dinh = lambda client, model, ts, phieu: ai.KiemDinh(diem=72, nhan_xet_chung="Cần xem lại câu 1", van_de=[
        ai.VanDe(cau_so=1, loai="dap_an_sai", mo_ta="Đáp án chưa đúng", de_xuat="Sửa thành B")])
    ai.tao_lai_cau = lambda client, model, ts, phieu, i, che_do, gop_y="": phieu.cau_hoi[i].model_copy(update={"noi_dung": gop_y})
    try:
        at = _app("fake")
        at.switch_page("trang/soan_phieu.py").run()
        at.button(key="mau_main").click().run()
        at.button(key="kd_lai").click().run()
        assert not at.exception
        assert at.session_state.ban.kiem_dinh.diem == 72
        assert any("Đáp án chưa đúng" in w.value for w in at.warning)
        at.button(key="gy0_0").click().run()
        assert not at.exception
        assert at.session_state.ban.phieu.cau_hoi[0].noi_dung.startswith("Đáp án chưa đúng")
        assert at.session_state.ban.kiem_dinh.van_de == []
    finally:
        ai.kiem_dinh, ai.tao_lai_cau = goc


def test_trang_lop_cua_toi():
    goc = ai.tao_de_b
    ai.tao_de_b = lambda client, model, ts, phieu: phieu.model_copy(deep=True)
    try:
        at = _app("fake")
        at.switch_page("trang/lop_cua_toi.py").run()
        at.text_area(key="dan_ten").input("Nguyễn An\nLê Bình\n\nTrần Chi").run()
        at.button(key="them_ten").click().run()
        assert not at.exception
        assert [h.ten for h in at.session_state.lop.hoc_sinh] == ["Nguyễn An", "Lê Bình", "Trần Chi"]
        at.run()  # chạy lại không được nhân đôi dòng
        assert len(at.session_state.lop.hoc_sinh) == 3
        at.button(key="mau_sb").click().run()
        at.button(key="tao_de_b").click().run()
        assert not at.exception and at.session_state.ban.phieu_b is not None
        at.toggle(key="xen_ke").set_value(True).run()
        assert not at.exception
    finally:
        ai.tao_de_b = goc


def test_trang_cham_bai_ket_qua():
    at = _app("fake")
    at.switch_page("trang/cham_bai.py").run()
    assert not at.exception  # chưa có phiếu: hướng dẫn, không lỗi
    at.button(key="mau_sb").click().run()
    at.session_state["lop"] = ai.LopHoc(ten_lop="2A", hoc_sinh=[ai.HocSinh(stt=1, ten="Nguyễn An", nhom="Cam")])
    at.session_state["cham"] = {"ver": 1, "loi": ["x.jpg: lỗi"], "rows": [
        {"STT": 1, "Họ tên": "Nguyễn An", "Điểm": 9.5, "Mức TT27": "Hoàn thành tốt", "Nhận xét": "Tốt", "Câu 1": "✓"},
        {"STT": 2, "Họ tên": "Lê Bình", "Điểm": 4.0, "Mức TT27": "Chưa hoàn thành", "Nhận xét": "Cố gắng", "Câu 1": "✗"}]}
    at.run()
    assert not at.exception
    assert [m.value for m in at.metric] == ["1 em", "0 em", "1 em"]
    at.button(key="cb_ap_nhom").click().run()
    assert {h.ten: h.nhom for h in at.session_state.lop.hoc_sinh} == {"Nguyễn An": "Tím", "Lê Bình": "Xanh"}


class DangNhap:
    """Giả lập giáo viên đã đăng nhập + database SQLite."""
    def __init__(self, email="co.lan@truong.vn", db=None):
        import state
        self.state, self.email, self.db = state, email, db or _sqlite()

    def __enter__(self):
        s = self.state
        self.goc = s.email, s.co_db, s.db
        s.email, s.co_db, s.db = (lambda: self.email), (lambda: True), (lambda: self.db)
        return self

    def __exit__(self, *a):
        self.state.email, self.state.co_db, self.state.db = self.goc


def test_dang_nhap_tu_luu_va_nap_ho_so():
    import kho
    with DangNhap() as dn:
        at = _app()
        at.text_input(key="ten_gv").input("Cô Lan").run()
        at.button(key="mau_sb").click().run()
        assert not at.exception
        ds = kho.phieu_ds(dn.db, dn.email)
        assert len(ds) == 1 and at.session_state.ban.id == ds[0]["id"]
        at.run()  # không đổi gì: không tạo thêm bản ghi
        assert len(kho.phieu_ds(dn.db, dn.email)) == 1
        assert kho.ho_so_lay(dn.db, dn.email)["ten"] == "Cô Lan"
        assert [m.value for m in at.metric][:2] == ["1", "0"]  # thống kê trang chủ
        kho.lop_luu(dn.db, dn.email, ai.LopHoc(ten_lop="2A", hoc_sinh=[ai.HocSinh(stt=1, ten="An")]))
        moi = _app()  # phiên mới: nạp hồ sơ + lớp gần nhất
        assert moi.session_state.ten_gv == "Cô Lan" and moi.session_state.lop.ten_lop == "2A"


def test_thu_vien():
    import kho
    with DangNhap() as dn:
        a = BAN.model_copy(deep=True)
        b = BAN.model_copy(deep=True)
        b.thong_so.chu_de, b.thong_so.mon = "Đọc hiểu: Chuyện quả bầu", "Tiếng Việt"
        kho.phieu_luu(dn.db, dn.email, a); kho.phieu_luu(dn.db, dn.email, b)
        kho.phieu_luu(dn.db, "nguoi.khac@x.vn", BAN.model_copy(deep=True))
        at = _app()
        at.switch_page("trang/thu_vien.py").run()
        dem = lambda: next(c.value for c in at.caption if c.value.startswith("Hiển thị"))
        assert not at.exception and "2/2" in dem()  # không thấy phiếu người khác
        at.text_input[0].input("qua bau").run()  # tìm không dấu
        assert "1/2" in dem()
        at.text_input[0].input("").run()
        at.button(key=f"tv_sao_{a.id}").click().run()
        assert kho.phieu_mo(dn.db, dn.email, a.id) and [p["gan_sao"] for p in kho.phieu_ds(dn.db, dn.email) if p["id"] == a.id] == [True]
        at.button(key=f"tv_nb_{a.id}").click().run()
        assert not at.exception and at.session_state.ban.thong_so.chu_de.endswith("(bản sao)")
        assert len(kho.phieu_ds(dn.db, dn.email)) == 3  # bản sao đã tự lưu
        at.switch_page("trang/thu_vien.py").run()
        at.button(key=f"tv_xoa_{b.id}").click().run()
        assert len(kho.phieu_ds(dn.db, dn.email)) == 2


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("OK", name)
