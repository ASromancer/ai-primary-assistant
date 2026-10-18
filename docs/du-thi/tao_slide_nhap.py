"""Tạo slide-nhap.pptx: bản nháp 15 trang (chữ + ảnh) để copy sang mẫu PowerPoint của Ban Tổ chức.

Chạy: .venv/bin/python docs/du-thi/tao_slide_nhap.py   (cần: pip install python-pptx)
Sửa nội dung ở danh sách TRANG bên dưới rồi chạy lại. Chỗ ⟦…⟧ là thông tin/số liệu thật cần điền.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

THU_MUC = Path(__file__).parent
ANH = THU_MUC / "anh"
XANH = RGBColor(0x0D, 0x47, 0xA1)
XAM = RGBColor(0x44, 0x44, 0x44)
FONT = "Be Vietnam Pro"

# (tiêu đề, [các dòng chữ], [ảnh])
TRANG = [
    ("TRỢ LÝ GIẢNG DẠY",
     ["Phiếu bài tập phân hóa và đánh giá học sinh tiểu học theo Thông tư 27 bằng AI",
      "⟦Họ tên tác giả⟧ – ⟦chức vụ⟧ – Trường ⟦…⟧, ⟦tỉnh/thành⟧",
      "Bảng A · Chủ đề: Giảng dạy, học tập và đánh giá"],
     ["01-trang-chu.png", "qr-app.png"]),
    ("Bối cảnh",
     ["Lớp ⟦…⟧: ⟦sĩ số⟧ học sinh, trình độ chênh lệch rõ",
      "TT27: đánh giá thường xuyên bằng nhận xét; bài tập 3 mức",
      "Phân hóa cần thiết nhưng tốn thời gian: soạn 3 mức, đáp án, in theo nhóm, chấm, nhận xét",
      "Soạn 1 phiếu ≈ ⟦…⟧ phút · chấm và nhận xét 1 lớp ≈ ⟦…⟧ phút"], []),
    ("Vấn đề và nguyên nhân",
     ["Soạn đề 3 mức đúng chuẩn TT27 tốn thời gian, chất lượng không đều",
      "Chatbot AI phổ thông: không theo mẫu TT27, không tự kiểm tra đáp án, không nối với chấm và theo dõi tiến bộ",
      "Kết quả chấm nằm trên giấy, khó phân tích để dạy bù kịp thời",
      "Làm sao để AI giúp đi hết vòng soạn → dạy → đánh giá → bổ trợ mà giáo viên vẫn làm chủ?"], []),
    ("Giải pháp: vòng khép kín 6 bước",
     ["① Soạn phiếu 3 mức → ② AI kiểm định → ③ Dạy và giao bài",
      "④ Chấm và nhận xét TT27 → ⑤ Phân tích lỗi và phiếu bổ trợ → ⑥ Tiến bộ và học bạ",
      "AI đề xuất – giáo viên quyết định"], ["so-do-vong-khep-kin.png"]),
    ("AI được dùng như thế nào",
     ["Soạn phiếu từ tên bài hoặc ảnh sách giáo khoa · AI kiểm định lượt 2 · chấm ảnh bài viết tay",
      "Nhận xét TT27 · phiếu bổ trợ từ phân tích lỗi · gợi ý nhận xét học bạ",
      "Công cụ: Google Gemini (đa phương thức) · Streamlit · Supabase · đăng nhập Google",
      "Chi phí: 0 đồng"], []),
    ("① Soạn phiếu từ tên bài hoặc ảnh sách giáo khoa",
     ["Phiếu 3 mức kèm đáp án, hướng dẫn chấm, ma trận, yêu cầu cần đạt",
      "Đáp án chuẩn tiểu học: Bài giải – lời giải – phép tính (đơn vị) – Đáp số"],
     ["02-soan-phieu-kiem-dinh.png"]),
    ("② AI kiểm định – tổ trưởng chuyên môn ảo",
     ["Tự giải lại từng câu: soát đáp án sai, câu lệch mức, từ ngữ chưa phù hợp",
      "Điểm chất lượng 0–100, góp ý từng câu, nút Sửa theo góp ý",
      "Thí điểm: phát hiện ⟦…⟧ lỗi / ⟦…⟧ phiếu"],
     ["03-gop-y-kiem-dinh.png"]),
    ("③ Phân hóa không dán nhãn",
     ["1 phiếu tự tách thành Phiếu Xanh / Cam / Tím (tên màu trung tính)",
      "Đề A/B chống nhìn bài · phiếu in sẵn tên từng em · in A4 hoặc Word"],
     ["31-phieu-xanh.png", "13-lop-cua-toi-phieu-theo-ten.png"]),
    ("③ Lớp học số: trình chiếu và giao bài bằng QR",
     ["Trình chiếu kiểu trò chơi: đếm giờ, lật đáp án, gọi tên ngẫu nhiên, đọc to",
      "Học sinh quét QR, chạm vào tên mình, không cần tài khoản, nhận đúng phiếu nhóm"],
     ["05-lop-hoc-trinh-chieu.png", "20-hs-chon-ten.png", "21-hs-lam-cau.png"]),
    ("④ Chấm và nhận xét theo TT27",
     ["Chấm từ ảnh bài viết tay hoặc bài online; trắc nghiệm, đúng/sai, nối cột chấm ngay",
      "Mức đạt theo quy tắc công khai: Chưa hoàn thành (M1 < 50%) · Hoàn thành tốt (M1, M2 ≥ 80%, M3 ≥ 50%) · còn lại Hoàn thành",
      "Giáo viên chấm tay hoặc sửa mọi kết quả"],
     ["09-cham-bai-so-tong-hop.png"]),
    ("⑤ Phân tích lỗi và phiếu bổ trợ",
     ["Tỉ lệ chưa đạt từng câu (ngưỡng 40%) và lỗi thường gặp",
      "1 nút: AI soạn phiếu bổ trợ đúng chỗ hổng kiến thức"],
     ["10-phan-tich-loi-sai.png"]),
    ("⑥ Theo dõi tiến bộ và nhận xét học bạ",
     ["Mức TT27 của từng em qua các lần đánh giá",
      "AI gợi ý nhận xét cuối kỳ (không gửi tên học sinh); giáo viên duyệt, xuất Excel"],
     ["12-tien-bo-bieu-do.png"]),
    ("Triển khai và kết quả thực tế",
     ["Phạm vi: lớp ⟦…⟧, ⟦…⟧ học sinh, ⟦…⟧ ngày, ⟦…⟧ đồng nghiệp dùng thử",
      "Soạn phiếu: ⟦…⟧ → ⟦…⟧ phút · chấm và nhận xét: ⟦…⟧ → ⟦…⟧ phút",
      "Câu yếu sau bổ trợ: chưa đạt ⟦…%⟧ → ⟦…%⟧ · học sinh hài lòng ⟦…%⟧",
      "\"⟦Trích phản hồi giáo viên⟧\""], []),
    ("AI có trách nhiệm · Bền vững và nhân rộng",
     ["AI đề xuất, giáo viên quyết định · kiểm định 2 lượt · không gửi tên học sinh cho AI",
      "Học sinh không cần tài khoản · dữ liệu từng giáo viên tách riêng",
      "0 đồng · chạy trên trình duyệt · khối 1–5, 8 môn, 3 bộ sách · nhiều giáo viên dùng song song",
      "Hướng phát triển: chia sẻ trong tổ chuyên môn, tin nhắn phụ huynh"], []),
    ("Kết luận",
     ["Để mỗi học sinh được học vừa sức, và giáo viên có thêm thời gian cho từng em",
      "Dùng thử: quét mã QR",
      "Phần lập trình có sự hỗ trợ của công cụ AI Claude Code ⟦và …⟧",
      "Trân trọng cảm ơn!"],
     ["qr-app.png"]),
]


def chu(khung, dong, co, mau=XAM, dam=False):
    p = khung.add_paragraph() if khung.paragraphs[0].text else khung.paragraphs[0]
    r = p.add_run()
    r.text, r.font.size, r.font.bold, r.font.name = dong, Pt(co), dam, FONT
    r.font.color.rgb = mau
    p.space_after = Pt(10)


def dat_anh(slide, ten, trai, tren, rong, cao):
    """Đặt ảnh vừa khít trong khung (giữ tỉ lệ)."""
    duong = ANH / ten
    if not duong.exists():
        return
    hinh = slide.shapes.add_picture(str(duong), trai, tren)
    ty_le = min(rong / hinh.width, cao / hinh.height)
    hinh.width, hinh.height = int(hinh.width * ty_le), int(hinh.height * ty_le)
    hinh.left = int(trai + (rong - hinh.width) / 2)
    hinh.top = int(tren + (cao - hinh.height) / 2)


def tao():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    for so, (tieu_de, dong, anh) in enumerate(TRANG, 1):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tt = s.shapes.add_textbox(Inches(0.6), Inches(0.35), W - Inches(1.2), Inches(1)).text_frame
        tt.word_wrap = True
        chu(tt, tieu_de, 34 if so > 1 else 44, XANH, True)
        co_anh = bool(anh)
        rong_chu = Inches(5.3) if co_anh else W - Inches(1.2)
        nd = s.shapes.add_textbox(Inches(0.6), Inches(1.6), rong_chu, H - Inches(2.2)).text_frame
        nd.word_wrap = True
        for d in dong:
            chu(nd, ("• " if so > 1 else "") + d, 20 if co_anh else 24)
        if co_anh:
            vung_trai, vung_rong = Inches(6.2), W - Inches(6.6)
            qr = [a for a in anh if a.startswith("qr")]
            anh_chinh = [a for a in anh if not a.startswith("qr")]
            if anh_chinh:
                cao = H - Inches(2.0)
                rong_moi = Emu(int(vung_rong / len(anh_chinh)))
                for k, a in enumerate(anh_chinh):
                    dat_anh(s, a, vung_trai + rong_moi * k, Inches(1.5), rong_moi - Inches(0.1), cao)
            for a in qr:
                dat_anh(s, a, W - Inches(2.6), H - Inches(2.6), Inches(2.2), Inches(2.2))
        so_trang = s.shapes.add_textbox(W - Inches(1.0), H - Inches(0.5), Inches(0.8), Inches(0.4)).text_frame
        chu(so_trang, f"{so}/{len(TRANG)}", 12)
    ra = THU_MUC / "slide-nhap.pptx"
    prs.save(ra)
    print(f"Đã tạo {ra} ({len(TRANG)} trang)")


if __name__ == "__main__":
    tao()
